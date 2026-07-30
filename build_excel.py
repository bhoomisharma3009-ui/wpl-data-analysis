import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.utils import get_column_letter

df = pd.read_csv("dataset.csv")
teams = sorted(set(df["team1"]) | set(df["team2"]))

wb = Workbook()

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
TITLE_FONT = Font(name="Arial", bold=True, size=14)
BASE_FONT = Font(name="Arial")

def style_header(ws, row, n_cols):
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

def autofit(ws, n_cols, width=20):
    for c in range(1, n_cols + 1):
        ws.column_dimensions[get_column_letter(c)].width = width

# ---------------------------------------------------------------
# Sheet 1: Raw Data
# ---------------------------------------------------------------
ws1 = wb.active
ws1.title = "Raw Data"
ws1["A1"] = "WPL Matches - Raw Data"
ws1["A1"].font = TITLE_FONT

headers = list(df.columns)
for c, h in enumerate(headers, start=1):
    ws1.cell(row=3, column=c, value=h)
style_header(ws1, 3, len(headers))

for r, row in enumerate(df.itertuples(index=False), start=4):
    for c, val in enumerate(row, start=1):
        ws1.cell(row=r, column=c, value=val)

autofit(ws1, len(headers), 16)
last_row = 3 + len(df)

# ---------------------------------------------------------------
# Sheet 2: Team Summary (formulas reference Raw Data)
# ---------------------------------------------------------------
ws2 = wb.create_sheet("Team Summary")
ws2["A1"] = "Team-wise Win Summary"
ws2["A1"].font = TITLE_FONT

ws2["A3"] = "Team"
ws2["B3"] = "Total Wins"
ws2["C3"] = "Matches Played"
ws2["D3"] = "Win %"
style_header(ws2, 3, 4)

for i, team in enumerate(teams, start=4):
    ws2.cell(row=i, column=1, value=team)
    # Total wins: COUNTIF on Raw Data winner column (column J = "winner")
    ws2.cell(row=i, column=2,
              value=f'=COUNTIF(\'Raw Data\'!$J$4:$J${last_row},A{i})')
    # Matches played: team1 OR team2 match
    ws2.cell(row=i, column=3,
              value=(f'=COUNTIF(\'Raw Data\'!$D$4:$D${last_row},A{i})'
                     f'+COUNTIF(\'Raw Data\'!$E$4:$E${last_row},A{i})'))
    ws2.cell(row=i, column=4, value=f'=IFERROR(B{i}/C{i},0)')
    ws2.cell(row=i, column=4).number_format = "0.0%"

autofit(ws2, 4, 22)

last_team_row = 3 + len(teams)

# Bar chart: Total wins per team
chart1 = BarChart()
chart1.title = "Total Wins by Team"
chart1.x_axis.title = "Team"
chart1.y_axis.title = "Wins"
data = Reference(ws2, min_col=2, min_row=3, max_row=last_team_row)
cats = Reference(ws2, min_col=1, min_row=4, max_row=last_team_row)
chart1.add_data(data, titles_from_data=True)
chart1.set_categories(cats)
chart1.width = 16
chart1.height = 9
ws2.add_chart(chart1, "F3")

# ---------------------------------------------------------------
# Sheet 3: Toss Analysis
# ---------------------------------------------------------------
ws3 = wb.create_sheet("Toss Analysis")
ws3["A1"] = "Toss Impact Analysis"
ws3["A1"].font = TITLE_FONT

ws3["A3"] = "Metric"
ws3["B3"] = "Count"
style_header(ws3, 3, 2)

ws3["A4"] = "Toss Winner Won Match"
# Row-wise comparison of toss_winner (H) vs winner (J)
ws3["B4"] = f"=SUMPRODUCT(--('Raw Data'!$H$4:$H${last_row}='Raw Data'!$J$4:$J${last_row}))"
ws3["A5"] = "Toss Winner Lost Match"
ws3["B5"] = f"=({last_row}-3)-B4"

ws3["A7"] = "Toss Decision"
ws3["B7"] = "Times Chosen"
style_header(ws3, 7, 2)
ws3["A8"] = "bat"
ws3["B8"] = f"=COUNTIF('Raw Data'!$I$4:$I${last_row},\"bat\")"
ws3["A9"] = "field"
ws3["B9"] = f"=COUNTIF('Raw Data'!$I$4:$I${last_row},\"field\")"

autofit(ws3, 2, 24)

# Pie chart: toss decision split
chart2 = PieChart()
chart2.title = "Toss Decision: Bat vs Field"
data2 = Reference(ws3, min_col=2, min_row=7, max_row=9)
cats2 = Reference(ws3, min_col=1, min_row=8, max_row=9)
chart2.add_data(data2, titles_from_data=True)
chart2.set_categories(cats2)
chart2.width = 12
chart2.height = 8
ws3.add_chart(chart2, "D3")

# Bar chart: toss winner won vs lost
chart3 = BarChart()
chart3.title = "Toss Winner Won vs Lost Match"
data3 = Reference(ws3, min_col=2, min_row=3, max_row=5)
cats3 = Reference(ws3, min_col=1, min_row=4, max_row=5)
chart3.add_data(data3, titles_from_data=True)
chart3.set_categories(cats3)
chart3.width = 12
chart3.height = 8
ws3.add_chart(chart3, "D14")

# ---------------------------------------------------------------
# Sheet 4: Season Trend
# ---------------------------------------------------------------
ws4 = wb.create_sheet("Season Trend")
ws4["A1"] = "Matches Played Per Season"
ws4["A1"].font = TITLE_FONT

seasons = sorted(df["season"].unique().tolist())
ws4["A3"] = "Season"
ws4["B3"] = "Matches Played"
style_header(ws4, 3, 2)

for i, s in enumerate(seasons, start=4):
    ws4.cell(row=i, column=1, value=s)
    ws4.cell(row=i, column=2,
              value=f'=COUNTIF(\'Raw Data\'!$B$4:$B${last_row},A{i})')

autofit(ws4, 2, 18)
last_season_row = 3 + len(seasons)

chart4 = BarChart()
chart4.title = "Matches Played Per Season"
chart4.x_axis.title = "Season"
chart4.y_axis.title = "Matches"
data4 = Reference(ws4, min_col=2, min_row=3, max_row=last_season_row)
cats4 = Reference(ws4, min_col=1, min_row=4, max_row=last_season_row)
chart4.add_data(data4, titles_from_data=True)
chart4.set_categories(cats4)
chart4.width = 16
chart4.height = 9
ws4.add_chart(chart4, "D3")

# ---------------------------------------------------------------
# Sheet 5: Venue Summary
# ---------------------------------------------------------------
ws5 = wb.create_sheet("Venue Summary")
ws5["A1"] = "Matches Hosted Per Venue"
ws5["A1"].font = TITLE_FONT

venues = sorted(df["venue"].unique().tolist())
ws5["A3"] = "Venue"
ws5["B3"] = "Matches Hosted"
style_header(ws5, 3, 2)

for i, v in enumerate(venues, start=4):
    ws5.cell(row=i, column=1, value=v)
    ws5.cell(row=i, column=2,
              value=f'=COUNTIF(\'Raw Data\'!$F$4:$F${last_row},A{i})')

autofit(ws5, 2, 26)
last_venue_row = 3 + len(venues)

chart5 = BarChart()
chart5.title = "Matches Hosted per Venue"
chart5.x_axis.title = "Venue"
chart5.y_axis.title = "Matches"
data5 = Reference(ws5, min_col=2, min_row=3, max_row=last_venue_row)
cats5 = Reference(ws5, min_col=1, min_row=4, max_row=last_venue_row)
chart5.add_data(data5, titles_from_data=True)
chart5.set_categories(cats5)
chart5.width = 16
chart5.height = 9
ws5.add_chart(chart5, "D3")

# ---------------------------------------------------------------
# Note sheet documenting data source assumption
# ---------------------------------------------------------------
ws6 = wb.create_sheet("Notes")
ws6["A1"] = "Data Source Notes"
ws6["A1"].font = TITLE_FONT
ws6["A3"] = ("This workbook uses a sample dataset with the same column structure as the "
             "real Kaggle WPL dataset (kaggle.com/datasets/sahiltailor/womens-premier-league-2023-2024-ball-by-ball). "
             "2023 & 2025 finals (Mumbai Indians) and 2024 final (RCB) reflect real results; "
             "other league-stage match outcomes are randomly generated for demo purposes. "
             "Replace 'Raw Data' with the real matches.csv (same headers) to make every "
             "number in this workbook reflect actual WPL results.")
ws6["A3"].alignment = Alignment(wrap_text=True)
ws6.column_dimensions["A"].width = 100
ws6.row_dimensions[3].height = 60

wb.save("WPL_Dashboard.xlsx")
print("WPL_Dashboard.xlsx created")
